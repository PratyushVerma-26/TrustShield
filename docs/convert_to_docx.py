import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import re
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """
    kwargs: top, bottom, left, right
    values: {"sz": 4, "val": "single", "color": "CCCCCC"}
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            b_xml = f'<w:{edge} {nsdecls("w")} w:val="{edge_data.get("val","single")}" w:sz="{edge_data.get("sz","4")}" w:space="0" w:color="{edge_data.get("color","D3D3D3")}"/>'
        else:
            b_xml = f'<w:{edge} {nsdecls("w")} w:val="none"/>'
        tcBorders.append(parse_xml(b_xml))
    tcPr.append(tcBorders)

def add_styled_code_block(doc, text_lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F4F6F9")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    set_cell_border(cell, 
                    top={"sz": 4, "val": "single", "color": "D0D7DE"},
                    bottom={"sz": 4, "val": "single", "color": "D0D7DE"},
                    left={"sz": 18, "val": "single", "color": "1B365D"}, # Accent left border
                    right={"sz": 4, "val": "single", "color": "D0D7DE"})
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run("\n".join(text_lines))
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)
    
    # Empty space after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(2)
    p_after.paragraph_format.space_after = Pt(6)

def add_header_footer(doc):
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("TrustShield: A Unified Digital Safety Platform for Multi-Threat Detection")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(0x70, 0x80, 0x90)
        hrun.italic = True
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun1 = fp.add_run("BCSE497J Capstone Project-I  |  School of Computer Science and Engineering (SCOPE)")
        frun1.font.name = "Calibri"
        frun1.font.size = Pt(8.5)
        frun1.font.color.rgb = RGBColor(0x70, 0x80, 0x90)

def parse_markdown_to_docx(md_path, docx_path):
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()
        
    doc = docx.Document()
    add_header_footer(doc)
    
    # Set default style to Times New Roman
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x20, 0x20, 0x20)
    
    i = 0
    n = len(lines)
    
    in_code_block = False
    code_lines = []
    
    while i < n:
        raw_line = lines[i]
        line = raw_line.strip()
        
        # Handle code blocks
        if line.startswith("```"):
            if in_code_block:
                add_styled_code_block(doc, code_lines)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue
            
        if in_code_block:
            code_lines.append(raw_line.rstrip("\r\n"))
            i += 1
            continue
            
        # Page break directive
        if "\\newpage" in line:
            doc.add_page_break()
            i += 1
            continue
            
        # Horizontal rule
        if line in ("---", "***", "___"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            # Add subtle line
            run = p.add_run("―" * 45)
            run.font.color.rgb = RGBColor(0xCC, 0xD1, 0xD9)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
            
        # Table detection
        if line.startswith("|") and line.endswith("|"):
            table_lines = []
            while i < n and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            
            # Parse table
            if len(table_lines) >= 2:
                rows_data = []
                for tl in table_lines:
                    # Skip separator line like | :--- | :--- |
                    if re.match(r'^\|(\s*:?-+:?\s*\|)+$', tl):
                        continue
                    cells = [c.strip() for c in tl.strip("|").split("|")]
                    rows_data.append(cells)
                
                if rows_data:
                    num_cols = max(len(r) for r in rows_data)
                    table = doc.add_table(rows=len(rows_data), cols=num_cols)
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    table.autofit = True
                    
                    for r_idx, row in enumerate(rows_data):
                        is_header = (r_idx == 0)
                        for c_idx in range(num_cols):
                            cell_text = row[c_idx] if c_idx < len(row) else ""
                            cell = table.cell(r_idx, c_idx)
                            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                            
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            p.paragraph_format.line_spacing = 1.05
                            
                            # Clean markdown bold/italics
                            clean_text = cell_text.replace("**", "").replace("*", "")
                            
                            run = p.add_run(clean_text)
                            run.font.name = "Calibri"
                            
                            if is_header:
                                set_cell_background(cell, "1B365D") # Deep Navy
                                run.font.size = Pt(9.5)
                                run.font.bold = True
                                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                                set_cell_border(cell, 
                                                top={"sz": 6, "val": "single", "color": "1B365D"},
                                                bottom={"sz": 12, "val": "single", "color": "0F2038"},
                                                left={"sz": 4, "val": "single", "color": "3A506B"},
                                                right={"sz": 4, "val": "single", "color": "3A506B"})
                            else:
                                bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                                set_cell_background(cell, bg_color)
                                run.font.size = Pt(9)
                                run.font.color.rgb = RGBColor(0x24, 0x29, 0x2F)
                                set_cell_border(cell, 
                                                top={"sz": 2, "val": "single", "color": "E2E8F0"},
                                                bottom={"sz": 2, "val": "single", "color": "E2E8F0"},
                                                left={"sz": 2, "val": "single", "color": "E2E8F0"},
                                                right={"sz": 2, "val": "single", "color": "E2E8F0"})
                    
                    p_spacer = doc.add_paragraph()
                    p_spacer.paragraph_format.space_after = Pt(4)
            continue
            
        # Headings
        if line.startswith("# ") or (line.startswith("# **") and line.endswith("**")):
            clean_title = re.sub(r'^[#\s\*]+|[#\s\*]+$', '', line)
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(clean_title)
            run.font.name = "Calibri"
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
            if "TRUSTSHIELD" in clean_title or "BCSE497J" in clean_title:
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
            
        if line.startswith("## ") or (line.startswith("## **") and line.endswith("**")):
            clean_title = re.sub(r'^[#\s\*]+|[#\s\*]+$', '', line)
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(clean_title)
            run.font.name = "Calibri"
            run.font.size = Pt(13.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2E, 0x40, 0x57)
            if any(k in clean_title for k in ["CERTIFICATE", "DECLARATION", "ACKNOWLEDGEMENTS", "ABSTRACT", "TABLE OF CONTENTS", "LIST OF"]):
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
            
        if line.startswith("### ") or (line.startswith("### **") and line.endswith("**")):
            clean_title = re.sub(r'^[#\s\*]+|[#\s\*]+$', '', line)
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(clean_title)
            run.font.name = "Calibri"
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            if any(k in clean_title for k in ["BACHELOR OF TECHNOLOGY", "COMPUTER SCIENCE", "Submitted By", "Under the Supervision"]):
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        if line.startswith("#### "):
            clean_title = re.sub(r'^[#\s\*]+|[#\s\*]+$', '', line)
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(clean_title)
            run.font.name = "Calibri"
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x4F, 0x5D, 0x75)
            i += 1
            continue
            
        # Bullet list items
        if line.startswith("- ") or line.startswith("* "):
            bullet_text = line[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            
            # Simple bold handler inside bullets: **bold**
            parts = re.split(r'(\*\*.*?\*\*)', bullet_text)
            for pt in parts:
                if pt.startswith("**") and pt.endswith("**"):
                    r = p.add_run(pt[2:-2])
                    r.bold = True
                else:
                    p.add_run(pt)
            i += 1
            continue
            
        # Numbered list items
        num_match = re.match(r'^(\d+)\.\s+(.*)$', line)
        if num_match:
            num_str = num_match.group(1)
            item_text = num_match.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            
            parts = re.split(r'(\*\*.*?\*\*)', item_text)
            for pt in parts:
                if pt.startswith("**") and pt.endswith("**"):
                    r = p.add_run(pt[2:-2])
                    r.bold = True
                else:
                    p.add_run(pt)
            i += 1
            continue

        # Blank lines
        if not line:
            i += 1
            continue

        # Regular paragraphs
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        
        # Center title page meta text
        if any(marker in line for marker in ["A Project Report Submitted", "ACADEMIC YEAR: 2025–2026", "September 2026", "Assistant Professor", "School of Computer Science and Engineering (SCOPE)"]):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("*Fig.") or line.startswith("Fig."):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)

        # Parse inline markdown (bold, italic, code)
        tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', line)
        for tok in tokens:
            if not tok:
                continue
            if tok.startswith("**") and tok.endswith("**"):
                r = p.add_run(tok[2:-2])
                r.bold = True
            elif tok.startswith("*") and tok.endswith("*"):
                r = p.add_run(tok[1:-1])
                r.italic = True
            elif tok.startswith("`") and tok.endswith("`"):
                r = p.add_run(tok[1:-1])
                r.font.name = "Consolas"
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0x80, 0x00, 0x20) # Burgundy code style
            else:
                p.add_run(tok)
                
        i += 1

    doc.save(docx_path)
    print(f"Document successfully created at {docx_path}")

if __name__ == "__main__":
    src_md = r"C:\Users\Pratyush\.gemini\antigravity\scratch\trustshield_docs\TRUSTSHIELD_FINAL_YEAR_PROJECT_REPORT.md"
    dst_docx = r"C:\Users\Pratyush\.gemini\antigravity\scratch\trustshield_docs\TrustShield_Final_Year_Project_Report.docx"
    parse_markdown_to_docx(src_md, dst_docx)
